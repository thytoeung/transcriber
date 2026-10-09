"""Write a TranscriptDoc to TXT / Markdown / DOCX / SRT / VTT."""

from __future__ import annotations

import re
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

from .document import Paragraph, TranscriptDoc, ends_sentence

FORMATS = {"txt": "Plain text", "md": "Markdown", "docx": "Word", "srt": "SRT subtitles",
           "vtt": "WebVTT subtitles"}
CUE_MAX_CHARS = 84
CUE_MAX_SEC = 6.0


@dataclass
class ExportOptions:
    timecodes: bool = True
    speakers: bool = True
    summary: bool = True


def fmt_time(sec: float) -> str:
    sec = max(0, int(sec))
    return f"{sec // 3600:02d}:{sec % 3600 // 60:02d}:{sec % 60:02d}"


def _fmt_cue_time(sec: float, sep: str) -> str:
    ms = max(0, int(round(sec * 1000)))
    return f"{ms // 3600000:02d}:{ms % 3600000 // 60000:02d}:{ms % 60000 // 1000:02d}{sep}{ms % 1000:03d}"


def output_path(doc_json: Path, fmt: str) -> Path:
    """'meeting.transcript.json' -> 'meeting.transcript.<fmt>'"""
    return doc_json.with_suffix(f".{fmt}")


def _prefix(doc: TranscriptDoc, p: Paragraph, opt: ExportOptions) -> str:
    parts = []
    if opt.timecodes:
        parts.append(f"[{fmt_time(p.start)}]")
    if opt.speakers and doc.diarized and p.speaker:
        parts.append(f"{doc.speaker_label(p.speaker)}:")
    return " ".join(parts)


def _header_lines(doc: TranscriptDoc) -> tuple[str, str]:
    title = Path(doc.audio_path).stem
    created = doc.created[:16].replace("T", " ")
    meta = f"Duration {fmt_time(doc.duration)} · Transcribed {created} · Model {doc.model}"
    if doc.diarized:
        names = ", ".join(doc.speaker_label(s) for s in doc.speaker_ids())
        meta += f" · Speakers: {names}"
    return title, meta


def to_txt(doc: TranscriptDoc, opt: ExportOptions) -> str:
    title, meta = _header_lines(doc)
    out = [title, meta, ""]
    if opt.summary and doc.summary.strip():
        out += ["SUMMARY", "=======", doc.summary.strip(), "", "TRANSCRIPT", "==========", ""]
    for p in doc.paragraphs:
        pre = _prefix(doc, p, opt)
        out += [f"{pre} {p.text}" if pre else p.text, ""]
    return "\n".join(out)


def to_markdown(doc: TranscriptDoc, opt: ExportOptions) -> str:
    title, meta = _header_lines(doc)
    out = [f"# {title}", "", f"*{meta}*", ""]
    if opt.summary and doc.summary.strip():
        out += ["## Summary", "", doc.summary.strip(), "", "## Transcript", ""]
    for p in doc.paragraphs:
        pre = _prefix(doc, p, opt)
        out += [f"**{pre}** {p.text}" if pre else p.text, ""]
    return "\n".join(out)


def _add_md_runs(par, text: str) -> None:
    for i, chunk in enumerate(re.split(r"\*\*(.+?)\*\*", text)):
        if chunk:
            par.add_run(chunk).bold = i % 2 == 1


def write_docx(doc: TranscriptDoc, opt: ExportOptions, path: Path) -> None:
    import docx
    from docx.shared import Pt, RGBColor

    d = docx.Document()
    d.styles["Normal"].font.name = "Calibri"
    d.styles["Normal"].font.size = Pt(11)
    title, meta = _header_lines(doc)
    d.add_heading(title, level=0)
    m = d.add_paragraph(meta)
    m.runs[0].italic = True

    if opt.summary and doc.summary.strip():
        d.add_heading("Summary", level=1)
        for line in doc.summary.strip().splitlines():
            s = line.strip()
            if not s:
                continue
            if h := re.match(r"^(#{1,4})\s+(.*)", s):
                d.add_heading(h.group(2).strip("* "), level=min(len(h.group(1)) + 1, 4))
            elif b := re.match(r"^[-*•]\s+(.*)", s):
                _add_md_runs(d.add_paragraph(style="List Bullet"), b.group(1))
            elif n := re.match(r"^\d+[.)]\s+(.*)", s):
                _add_md_runs(d.add_paragraph(style="List Number"), n.group(1))
            else:
                _add_md_runs(d.add_paragraph(), s)
        d.add_heading("Transcript", level=1)

    for p in doc.paragraphs:
        par = d.add_paragraph()
        par.paragraph_format.space_after = Pt(8)
        if opt.timecodes:
            r = par.add_run(f"[{fmt_time(p.start)}] ")
            r.font.color.rgb = RGBColor(0x80, 0x80, 0x80)
        if opt.speakers and doc.diarized and p.speaker:
            par.add_run(f"{doc.speaker_label(p.speaker)}: ").bold = True
        par.add_run(p.text)
    d.save(str(path))


# ----- subtitles -----

def _cues_for(p: Paragraph) -> list[tuple[float, float, str]]:
    """Short subtitle cues. Uses word timings unless the paragraph text was hand-edited."""
    cues = []
    if p.words and not p.edited:
        cur = []
        for w in p.words:
            cur.append(w)
            text = " ".join(x.text for x in cur)
            if (len(text) >= CUE_MAX_CHARS or cur[-1].end - cur[0].start >= CUE_MAX_SEC
                    or (ends_sentence(w.text) and len(text) > CUE_MAX_CHARS * 0.5)):
                cues.append((cur[0].start, cur[-1].end, text))
                cur = []
        if cur:
            cues.append((cur[0].start, cur[-1].end, " ".join(x.text for x in cur)))
        return cues
    # Edited: split text into chunks and spread them across the paragraph's time span.
    chunks, cur = [], ""
    for tok in p.text.split():
        if cur and len(cur) + 1 + len(tok) > CUE_MAX_CHARS:
            chunks.append(cur)
            cur = tok
        else:
            cur = f"{cur} {tok}".strip()
    if cur:
        chunks.append(cur)
    total = sum(len(c) for c in chunks) or 1
    t = p.start
    for c in chunks:
        dur = (p.end - p.start) * len(c) / total
        cues.append((t, t + dur, c))
        t += dur
    return cues


def _subtitle_cues(doc: TranscriptDoc, opt: ExportOptions):
    for p in doc.paragraphs:
        label = doc.speaker_label(p.speaker) if (opt.speakers and doc.diarized and p.speaker) else ""
        for i, (s, e, text) in enumerate(_cues_for(p)):
            yield s, max(e, s + 0.5), (f"{label}: {text}" if label and i == 0 else text), label


def to_srt(doc: TranscriptDoc, opt: ExportOptions) -> str:
    out = []
    for n, (s, e, text, _) in enumerate(_subtitle_cues(doc, opt), 1):
        out += [str(n), f"{_fmt_cue_time(s, ',')} --> {_fmt_cue_time(e, ',')}", text, ""]
    return "\n".join(out)


def to_vtt(doc: TranscriptDoc, opt: ExportOptions) -> str:
    out = ["WEBVTT", ""]
    for s, e, text, label in _subtitle_cues(doc, opt):
        if label and text.startswith(f"{label}: "):
            text = text[len(label) + 2:]
        out += [f"{_fmt_cue_time(s, '.')} --> {_fmt_cue_time(e, '.')}",
                f"<v {label}>{text}" if label else text, ""]
    return "\n".join(out)


def export(doc: TranscriptDoc, doc_json: Path, formats: list[str], opt: ExportOptions) -> list[Path]:
    written = []
    for fmt in formats:
        path = output_path(Path(doc_json), fmt)
        if fmt == "docx":
            write_docx(doc, opt, path)
        else:
            text = {"txt": to_txt, "md": to_markdown, "srt": to_srt, "vtt": to_vtt}[fmt](doc, opt)
            path.write_text(text, encoding="utf-8")
        written.append(path)
    return written


def now_iso() -> str:
    return datetime.now().isoformat(timespec="seconds")
